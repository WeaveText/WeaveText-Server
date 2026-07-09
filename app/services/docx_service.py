"""DOCX 处理服务。"""

import re
from pathlib import Path

from docx import Document
from docxtpl import DocxTemplate


class DocxService:
    """处理 DOCX 解析与渲染的服务层。"""

    _PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*([a-zA-Z_][a-zA-Z0-9_\.]*)\s*\}\}")

    def extract_placeholders_from_docx(self, docx_path: Path) -> list[str]:
        """
        从 DOCX 模板中提取 `{{placeholder}}` 占位符列表。
        """
        document = Document(str(docx_path))
        extracted: set[str] = set()

        for paragraph in document.paragraphs:
            extracted.update(self._PLACEHOLDER_PATTERN.findall(paragraph.text))

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    extracted.update(self._PLACEHOLDER_PATTERN.findall(cell.text))

        return sorted(extracted)

    def generate_placeholder_descriptions(self, placeholders: list[str]) -> dict[str, str]:
        """
        为占位符生成基础描述（当前为 Mock）。
        """
        return {
            placeholder: f"请提供与“{placeholder}”相关的内容，用于生成文档。"
            for placeholder in placeholders
        }

    def render_docx_template(
        self,
        template_path: Path,
        data: dict[str, object],
        output_path: Path,
    ) -> None:
        """
        使用 docxtpl 渲染模板并输出文件。
        """
        tpl = DocxTemplate(str(template_path))
        tpl.render(data)
        tpl.save(str(output_path))
